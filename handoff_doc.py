"""Volunteer Handoff Guide: one Markdown file shown on a Help page and downloaded as a Word document.

Sources: md/HANDOFF_GUIDE.md (public guide) and md/HANDOFF_PRIVATE_TEMPLATE.md (blank Contacts & Links
sheet, filled in on the PTA Google Drive - never in the repo).

Only the Markdown the guide uses is supported: # headings, paragraphs, "- " / "1. " lists nested by
2 spaces ("- [ ] " = checkbox), | tables |, ``` code blocks, "> " notes, **bold**, *italic*, `code`,
[links](url), bare http(s) URLs, and [TBD: what to look up] placeholders. Placeholders are highlighted
and listed at the end under "Still to look up".
"""

import html
import io
import re
from typing import Any, Dict, List, Optional

GUIDE_PATH = 'md/HANDOFF_GUIDE.md'
PRIVATE_TEMPLATE_PATH = 'md/HANDOFF_PRIVATE_TEMPLATE.md'
STILL_TO_LOOK_UP = 'Still to look up'

INLINE_RE = re.compile(
    r'(?P<tbd>\[TBD(?::\s*(?P<tbd_text>[^\]]*))?\])'
    r'|(?P<link>\[(?P<link_text>[^\]]+)\]\((?P<url>[^)\s]+)\))'
    r'|(?P<code>`(?P<code_text>[^`]+)`)'
    r'|(?P<bold>\*\*(?P<bold_text>.+?)\*\*)'
    r'|(?P<italic>\*(?P<italic_text>[^*\s][^*]*)\*)'
    r'|(?P<autolink>https?://[^\s)|]+[^\s)|.,;:])'
)
LIST_RE = re.compile(r'^(?P<indent> *)(?P<marker>[-*]|\d+\.) +(?P<text>.*)$')
HEADING_RE = re.compile(r'^(?P<hashes>#{1,4}) +(?P<text>.+)$')


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------

def parse_inline(text: str) -> List[Dict[str, str]]:
    """'a **b** [TBD: c]' -> [{'kind': 'text', 'text': 'a '}, {'kind': 'bold', ...}, {'kind': 'tbd', 'text': 'c'}]"""
    spans, pos = [], 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            spans.append({'kind': 'text', 'text': text[pos:m.start()]})
        if m.group('tbd'):
            spans.append({'kind': 'tbd', 'text': (m.group('tbd_text') or '').strip()})
        elif m.group('link'):
            spans.append({'kind': 'link', 'text': m.group('link_text'), 'url': m.group('url')})
        elif m.group('code'):
            spans.append({'kind': 'code', 'text': m.group('code_text')})
        elif m.group('bold'):
            spans.append({'kind': 'bold', 'text': m.group('bold_text')})
        elif m.group('italic'):
            spans.append({'kind': 'italic', 'text': m.group('italic_text')})
        else:
            spans.append({'kind': 'link', 'text': m.group('autolink'), 'url': m.group('autolink')})
        pos = m.end()
    if pos < len(text):
        spans.append({'kind': 'text', 'text': text[pos:]})
    return spans


def plain_text(text: str) -> str:
    """Inline markup removed, TBD placeholders dropped"""
    return ''.join(s['text'] for s in parse_inline(text) if s['kind'] != 'tbd').strip(' :-')


def slugify(text: str) -> str:
    return re.sub(r'[^a-z0-9]+', '-', plain_text(text).lower()).strip('-')


def _table_cells(line: str) -> List[str]:
    return [c.strip() for c in line.strip().strip('|').split('|')]


def parse_blocks(markdown: str) -> List[Dict[str, Any]]:
    """Markdown -> blocks: heading, para, quote, list, table, code"""
    blocks: List[Dict[str, Any]] = []
    lines = markdown.splitlines()
    para: List[str] = []

    def flush_para():
        if para:
            blocks.append({'type': 'para', 'text': ' '.join(para)})
            para.clear()

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            flush_para()
            i += 1
        elif stripped.startswith('```'):
            flush_para()
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code.append(lines[i])
                i += 1
            blocks.append({'type': 'code', 'text': '\n'.join(code)})
            i += 1
        elif HEADING_RE.match(stripped):
            flush_para()
            m = HEADING_RE.match(stripped)
            blocks.append({'type': 'heading', 'level': len(m.group('hashes')), 'text': m.group('text'),
                           'id': slugify(m.group('text'))})
            i += 1
        elif stripped.startswith('|'):
            flush_para()
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = _table_cells(lines[i])
                if not all(re.fullmatch(r':?-+:?', c) for c in cells):  # skip the |---| separator row
                    rows.append(cells)
                i += 1
            blocks.append({'type': 'table', 'header': rows[0], 'rows': rows[1:]})
        elif stripped.startswith('>'):
            flush_para()
            quote = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                quote.append(lines[i].strip()[1:].strip())
                i += 1
            blocks.append({'type': 'quote', 'text': ' '.join(quote)})
        elif LIST_RE.match(line):
            flush_para()
            items = []
            while i < len(lines) and LIST_RE.match(lines[i]):
                m = LIST_RE.match(lines[i])
                text = m.group('text')
                checkbox = text.startswith('[ ] ')
                items.append({'level': len(m.group('indent')) // 2, 'ordered': m.group('marker')[0].isdigit(),
                              'checkbox': checkbox, 'text': text[4:] if checkbox else text})
                i += 1
            blocks.append({'type': 'list', 'items': items})
        else:
            para.append(stripped)
            i += 1
    flush_para()
    return blocks


def _item_context(text: str) -> str:
    """A list item's bold lead-in ('**Secretary:** ...' -> 'Secretary'), else its text"""
    spans = [s for s in parse_inline(text) if s['text'].strip()]
    if spans and spans[0]['kind'] == 'bold':
        return spans[0]['text'].strip(' :.')
    return plain_text(text).rstrip('.')


def collect_tbds(blocks: List[Dict[str, Any]]) -> List[Dict[str, str]]:
    """Every [TBD] placeholder with where it is: [{'section', 'context', 'text'}].

    Bare [TBD] cells in a table are grouped per column ('This year: Kickoff assembly, Final assembly').
    """
    found = []
    section = ''

    def add(text, context):
        for span in parse_inline(text):
            if span['kind'] == 'tbd':
                found.append({'section': section, 'context': context, 'text': span['text']})

    for block in blocks:
        if block['type'] == 'heading':
            if block['level'] <= 3:
                section = plain_text(block['text'])
            add(block['text'], plain_text(block['text']))
        elif block['type'] in ('para', 'quote'):
            add(block['text'], '')
        elif block['type'] == 'list':
            for item in block['items']:
                add(item['text'], _item_context(item['text']))
        elif block['type'] == 'table':
            bare: Dict[str, List[str]] = {}  # column header -> row labels with a bare [TBD]
            for row in block['rows']:
                label = plain_text(row[0])
                for col, (header, cell) in enumerate(zip(block['header'], row)):
                    if col and re.fullmatch(r'\[TBD\]', cell):
                        bare.setdefault(plain_text(header), []).append(label)
                    else:
                        add(cell, f"{label} ({plain_text(header)})" if col else label)
            for header, labels in bare.items():
                found.append({'section': section, 'context': header, 'text': '; '.join(labels)})
    return found


def tbd_label(tbd: Dict[str, str]) -> str:
    """'Key dates › Kickoff assembly (This year)' plus ': detail' when the placeholder says what to find"""
    context = tbd['context'] if tbd['context'] != tbd['section'] else ''
    if len(context) > 70:
        context = context[:67].rstrip() + '...'
    where = ' › '.join(p for p in (tbd['section'], context) if p)
    return f"{where}: {tbd['text']}" if tbd['text'] else where


def load(path: str) -> Dict[str, Any]:
    """Parse a guide file: title (the # heading), blocks after it, and its TBD placeholders"""
    with open(path, encoding='utf-8') as f:
        blocks = parse_blocks(f.read())
    title = ''
    if blocks and blocks[0]['type'] == 'heading' and blocks[0]['level'] == 1:
        title = plain_text(blocks.pop(0)['text'])
    tbds = collect_tbds(blocks)
    if tbds:
        blocks = blocks + [{'type': 'heading', 'level': 2, 'text': STILL_TO_LOOK_UP, 'id': slugify(STILL_TO_LOOK_UP)},
                           {'type': 'list', 'items': [{'level': 0, 'ordered': False, 'checkbox': True,
                                                       'text': _escape_markup(tbd_label(t))} for t in tbds]}]
    return {'title': title, 'blocks': blocks, 'tbds': tbds,
            'toc': [{'id': b['id'], 'text': plain_text(b['text'])} for b in blocks
                    if b['type'] == 'heading' and b['level'] == 2]}


def _escape_markup(text: str) -> str:
    """Keep generated text (TBD labels) from being read as markup again"""
    return text.replace('[', '(').replace(']', ')').replace('*', '').replace('`', "'")


# ---------------------------------------------------------------------------
# HTML (Help page)
# ---------------------------------------------------------------------------

def inline_html(text: str) -> str:
    out = []
    for s in parse_inline(text):
        t = html.escape(s['text'])
        if s['kind'] == 'tbd':
            out.append(f'<mark class="tbd">⚠ TBD{": " + t if t else ""}</mark>')
        elif s['kind'] == 'link':
            out.append(f'<a href="{html.escape(s["url"])}" target="_blank" rel="noopener">{t}</a>')
        elif s['kind'] == 'code':
            out.append(f'<code>{t}</code>')
        elif s['kind'] == 'bold':
            out.append(f'<strong>{t}</strong>')
        elif s['kind'] == 'italic':
            out.append(f'<em>{t}</em>')
        else:
            out.append(t)
    return ''.join(out)


def _list_html(items: List[Dict[str, Any]]) -> str:
    """Nested <ul>/<ol> from items that carry a nesting level"""
    out, stack = [], []  # stack of open list tags; the last <li> is left open for a nested list
    for item in items:
        depth = item['level'] + 1
        if depth > len(stack):
            tag = 'ol' if item['ordered'] else 'ul'
            css = ' class="handoff-checklist"' if item['checkbox'] else ''
            while len(stack) < depth:
                out.append(f'<{tag}{css}>')
                stack.append(tag)
        else:
            out.append('</li>')
            while len(stack) > depth:
                out.append(f'</{stack.pop()}></li>')
        box = '<span class="handoff-box">☐</span> ' if item['checkbox'] else ''
        out.append(f'<li>{box}{inline_html(item["text"])}')
    if stack:
        out.append('</li>')
    while stack:
        out.append(f'</{stack.pop()}>')
        if stack:
            out.append('</li>')
    return ''.join(out)


def to_html(blocks: List[Dict[str, Any]]) -> str:
    out = []
    for b in blocks:
        if b['type'] == 'heading':
            tag = f"h{min(b['level'] + 1, 6)}"  # the page title is the h1, so ## -> h3
            out.append(f'<{tag} id="{b["id"]}" class="handoff-h{b["level"]}">{inline_html(b["text"])}</{tag}>')
        elif b['type'] == 'para':
            out.append(f'<p>{inline_html(b["text"])}</p>')
        elif b['type'] == 'quote':
            out.append(f'<div class="alert alert-info handoff-note">{inline_html(b["text"])}</div>')
        elif b['type'] == 'list':
            out.append(_list_html(b['items']))
        elif b['type'] == 'table':
            head = ''.join(f'<th>{inline_html(c)}</th>' for c in b['header'])
            body = ''.join('<tr>' + ''.join(f'<td>{inline_html(c)}</td>' for c in row) + '</tr>' for row in b['rows'])
            out.append(f'<div class="table-responsive"><table class="table table-sm table-bordered handoff-table">'
                       f'<thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>')
        elif b['type'] == 'code':
            out.append('<div class="handoff-code"><button type="button" class="btn btn-sm btn-outline-secondary '
                       'handoff-copy" onclick="copyHandoffCode(this)"><i class="bi bi-clipboard"></i> Copy</button>'
                       f'<pre><code>{html.escape(b["text"])}</code></pre></div>')
    return '\n'.join(out)


# ---------------------------------------------------------------------------
# Word (.docx) - needs python-docx (requirements.txt); imported only when a download is asked for
# ---------------------------------------------------------------------------

def docx_available() -> bool:
    try:
        import docx  # noqa: F401
        return True
    except ImportError:
        return False


TBD_FILL = 'FFF3CD'      # same pale yellow as the Help page highlight
HEADER_FILL = 'DBEAFE'   # table header, the app's soft blue (card-header-blue)
CODE_FILL = 'F3F4F6'
TEXT_WIDTH = 6.9 * 914400  # EMU: 8.5in page - 2 x 0.8in margins (to_docx)
TABLE_INDENT = 0.3 * 914400  # EMU: tables sit in from the text, level with the bullet text
# w:pPr children that must come after w:shd (OOXML requires this order)
_PPR_AFTER_SHD = ('w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap', 'w:overflowPunct',
                  'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN', 'w:bidi', 'w:adjustRightInd',
                  'w:snapToGrid', 'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents',
                  'w:suppressOverlap', 'w:jc', 'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap',
                  'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange')


def _shading(fill: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    return shd


def _add_hyperlink(paragraph, url: str, text: str):
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    style = OxmlElement('w:rStyle')
    style.set(qn('w:val'), 'Hyperlink')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    props.extend([style, color, underline])
    run.append(props)
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _add_inline(paragraph, text: str, bold: bool = False):
    from docx.enum.text import WD_COLOR_INDEX
    for s in parse_inline(text):
        if s['kind'] == 'link':
            _add_hyperlink(paragraph, s['url'], s['text'])
            continue
        if s['kind'] == 'tbd':
            run = paragraph.add_run(f"⚠ TBD{': ' + s['text'] if s['text'] else ''}")
            run.bold = True
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            continue
        run = paragraph.add_run(s['text'])
        run.bold = bold or s['kind'] == 'bold' or None
        run.italic = s['kind'] == 'italic' or None
        if s['kind'] == 'code':
            run.font.name = 'Consolas'


def _hanging(paragraph, level: int):
    from docx.shared import Inches
    paragraph.paragraph_format.left_indent = Inches(0.3 * (level + 1))
    paragraph.paragraph_format.first_line_indent = Inches(-0.22)


def _add_list(doc, items: List[Dict[str, Any]]):
    counters: Dict[int, int] = {}
    for item in items:
        level = min(item['level'], 2)
        for deeper in [k for k in counters if k > level]:
            del counters[deeper]
        if item['checkbox'] or item['ordered']:
            # Word's built-in numbering continues across lists, so number (and tick) by hand
            counters[level] = counters.get(level, 0) + 1
            p = doc.add_paragraph()
            _hanging(p, level)
            p.add_run('☐ ' if item['checkbox'] else f"{counters[level]}. ")
        else:
            p = doc.add_paragraph(style='List Bullet' if level == 0 else f'List Bullet {level + 1}')
            _hanging(p, level)  # explicit, so viewers that ignore list styles still indent
        p.paragraph_format.space_after = 0
        _add_inline(p, item['text'])


def _add_borders(table):
    """Explicit cell borders (some viewers ignore the Table Grid style)"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), 'BFC5CD')
        borders.append(el)
    tbl_pr = table._tbl.tblPr
    tbl_pr.insert_element_before(borders, 'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook',
                                 'w:tblCaption', 'w:tblDescription', 'w:tblPrChange')


def _add_table(doc, block: Dict[str, Any]):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    table = doc.add_table(rows=1 + len(block['rows']), cols=len(block['header']))
    table.style = 'Table Grid'
    _add_borders(table)
    # Indented from the text, filling the rest of the line; columns sized by their longest text (at least
    # 12 characters, so blank fill-in columns stay usable) - without this Word shrinks empty columns to nothing
    width_left = TEXT_WIDTH - TABLE_INDENT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn('w:tblW'))
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_w.set(qn('w:w'), str(int(width_left / 635)))  # EMU -> twentieths of a point
    tbl_ind = OxmlElement('w:tblInd')
    tbl_ind.set(qn('w:type'), 'dxa')
    tbl_ind.set(qn('w:w'), str(int(TABLE_INDENT / 635)))
    tbl_pr.insert_element_before(tbl_ind, 'w:tblBorders', 'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook',
                                 'w:tblCaption', 'w:tblDescription', 'w:tblPrChange')
    weights = [min(60, max([12] + [len(plain_text(row[c])) for row in [block['header']] + block['rows'] if c < len(row)]))
               for c in range(len(block['header']))]
    for c, weight in enumerate(weights):
        width = int(width_left * weight / sum(weights))
        table.columns[c].width = width  # the grid, read by some viewers
        for cell in table.columns[c].cells:
            cell.width = width          # the cells, read by Word
    for r, row in enumerate([block['header']] + block['rows']):
        for c, text in enumerate(row[:len(block['header'])]):
            cell = table.cell(r, c)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = 0
            _add_inline(p, text, bold=(r == 0))
            for run in p.runs:
                run.font.size = Pt(9.5)
            if r == 0:
                cell._tc.get_or_add_tcPr().append(_shading(HEADER_FILL))
            elif '[TBD' in text:
                cell._tc.get_or_add_tcPr().append(_shading(TBD_FILL))
    doc.add_paragraph().paragraph_format.space_after = 0


def _add_code(doc, text: str):
    from docx.shared import Inches, Pt
    p = doc.add_paragraph()
    p._p.get_or_add_pPr().insert_element_before(_shading(CODE_FILL), *_PPR_AFTER_SHD)
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.right_indent = Inches(0.15)
    for n, line in enumerate(text.split('\n')):
        run = p.add_run(line)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        if n < len(text.split('\n')) - 1:
            run.add_break()


def to_docx(guide: Dict[str, Any], footer: str = '') -> bytes:
    """A loaded guide (see load()) as a Word document"""
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.8)
        section.top_margin = section.bottom_margin = Inches(0.7)
        if footer:
            fp = section.footer.paragraphs[0]
            run = fp.add_run(footer)
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)

    doc.add_heading(guide['title'], 0)
    if guide['tbds']:
        p = doc.add_paragraph()
        run = p.add_run(f"⚠ Draft: {len(guide['tbds'])} items still to look up "
                        f"(listed at the end under \"{STILL_TO_LOOK_UP}\").")
        run.bold = True
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    for b in guide['blocks']:
        if b['type'] == 'heading':
            _add_inline(doc.add_heading(level=min(b['level'] - 1, 3)), b['text'])
        elif b['type'] == 'para':
            _add_inline(doc.add_paragraph(), b['text'])
        elif b['type'] == 'quote':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            _add_inline(p, b['text'])
            for run in p.runs:
                run.italic = True
        elif b['type'] == 'list':
            _add_list(doc, b['items'])
        elif b['type'] == 'table':
            _add_table(doc, b)
        elif b['type'] == 'code':
            _add_code(doc, b['text'])

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
